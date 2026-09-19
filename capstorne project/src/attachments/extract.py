"""Turn an uploaded PDF, Word document, or image into plain text for the LLM."""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass

MAX_CHARS = 10_000
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
PDF_SUFFIXES = {".pdf"}
WORD_SUFFIXES = {".docx"}
OLD_WORD_SUFFIXES = {".doc"}


@dataclass
class ExtractedFile:
    name: str
    kind: str
    text: str
    error: str | None = None

    @property
    def usable(self) -> bool:
        return bool((self.text or "").strip()) and self.error is None


def extract_upload(name: str, data: bytes, mime: str | None = None) -> ExtractedFile:
    """Read one upload. Images go through a vision model when a Groq key is set."""
    filename = (name or "attachment").strip() or "attachment"
    suffix = _suffix(filename)
    if not data:
        return ExtractedFile(filename, "empty", "", "The file was empty.")

    if suffix in PDF_SUFFIXES or (mime or "").endswith("pdf"):
        return _from_pdf(filename, data)
    if suffix in WORD_SUFFIXES or "wordprocessingml" in (mime or ""):
        return _from_docx(filename, data)
    if suffix in OLD_WORD_SUFFIXES:
        return ExtractedFile(
            filename,
            "word",
            "",
            "Older .doc files are not supported. Save the document as .docx or PDF.",
        )
    if suffix in IMAGE_SUFFIXES or (mime or "").startswith("image/"):
        return _from_image(filename, data, mime or _image_mime(suffix))
    return ExtractedFile(
        filename,
        "other",
        "",
        "Use a PDF, Word (.docx), or image file (PNG, JPG, WEBP).",
    )


def _suffix(name: str) -> str:
    return "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _image_mime(suffix: str) -> str:
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(suffix, "image/png")


def _clip(text: str) -> str:
    cleaned = " ".join((text or "").split())
    if len(cleaned) <= MAX_CHARS:
        return cleaned
    return cleaned[:MAX_CHARS] + "…"


def _from_pdf(name: str, data: bytes) -> ExtractedFile:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ExtractedFile(name, "pdf", "", "PDF support is not installed (pypdf).")
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages[:20]:
            pages.append(page.extract_text() or "")
        text = _clip("\n".join(pages))
        if not text:
            return ExtractedFile(
                name,
                "pdf",
                "",
                "No readable text was found in this PDF. If it is a scan, attach it as an image.",
            )
        return ExtractedFile(name, "pdf", text)
    except Exception as exc:  # noqa: BLE001 - surfaced to the chat, not swallowed
        return ExtractedFile(name, "pdf", "", f"Could not read this PDF: {exc}")


def _from_docx(name: str, data: bytes) -> ExtractedFile:
    try:
        from docx import Document
    except ImportError:
        return ExtractedFile(name, "word", "", "Word support is not installed (python-docx).")
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = _clip("\n".join(parts))
        if not text:
            return ExtractedFile(name, "word", "", "No readable text was found in this document.")
        return ExtractedFile(name, "word", text)
    except Exception as exc:  # noqa: BLE001
        return ExtractedFile(name, "word", "", f"Could not read this Word file: {exc}")


def _from_image(name: str, data: bytes, mime: str) -> ExtractedFile:
    """Ask Groq's vision model to transcribe the image; fall back to a clear error."""
    from src.config import settings

    if not settings.llm_ready:
        return ExtractedFile(
            name,
            "image",
            "",
            "Images need a Groq API key so the model can read them. Type the text, or attach a PDF/Word file.",
        )
    try:
        from langchain_core.messages import HumanMessage

        from src.agent.llm import get_llm

        encoded = base64.b64encode(data).decode("ascii")
        prompt = (
            "Transcribe every readable word from this image verbatim. "
            "Preserve labels such as Patient Name, Diagnosis, Medications. "
            "If it is a medical form or report, keep the field names. "
            "Reply with the text only."
        )
        model = get_llm(temperature=0, model=settings.groq_vision_model, max_tokens=1200)
        reply = model.invoke(
            [
                HumanMessage(
                    content=[
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:{mime};base64,{encoded}"},
                        },
                    ]
                )
            ]
        )
        text = _clip(str(getattr(reply, "content", "") or ""))
        if not text:
            return ExtractedFile(name, "image", "", "The vision model returned no text from this image.")
        return ExtractedFile(name, "image", text)
    except Exception as exc:  # noqa: BLE001
        return ExtractedFile(
            name,
            "image",
            "",
            f"Could not read this image ({exc}). Type the patient name or paste the text.",
        )
